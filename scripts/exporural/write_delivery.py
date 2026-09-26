"""Generate the same populated content in Markdown and Word. Offline only."""
import json, subprocess, hashlib, re
from pathlib import Path
from decimal import Decimal
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path('docs/exporural/2028-revisao-2026-09-25')
read=lambda f:json.loads((root/f).read_text(encoding='utf-8-sig'))
meta=read('fontes_e_revisao.json'); cal=read('calibracao.json')
manifest=read('manifesto_lotes.json'); cross=read('crosswalk_linhagem.json'); roads=read('geometria_vias.json')['roads']
commit=subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip()
revision=meta['revision']
blocks=[]
def add(kind,text):blocks.append((kind,text))
def table(headers,rows):blocks.append(('table',[headers,*rows]))
def fmt(v):return f'{Decimal(str(v)):,.2f}'.replace(',','X').replace('.',',').replace('X','.')
def codes(items):
    # Compact only consecutive codes within the SAME explicit lineage group.
    output=[];start=end=None
    for code in items:
        if end and code[:4]==end[:4] and int(code[4:])==int(end[4:])+1:end=code;continue
        if start:output.append(start if start==end else start+' a '+end)
        start=end=code
    if start:output.append(start if start==end else start+' a '+end)
    return ', '.join(output)

add('title','Revisão cadastral e geométrica da Exporural FENASOJA 2028')
add('p','Relatório de entrega do Codex para Leonardo e equipe Lovable. Data de referência: 25 de setembro de 2026. Este pacote implementa uma proposta local de 100 parcelas e a via transversal, preserva a autoridade dos dados persistidos e prepara a revisão de banco. O preview está funcional, mas a aceitação geométrica está bloqueada pelas diferenças de calibração/área e pelos apoios B37/B38 sobrepostos aos lotes S. A identidade física dos registros e a aplicação ao banco também continuam pendentes. Esta geometria ainda não está liberada para migração.')
add('p','Nenhuma alteração foi aplicada ao Supabase pelo Codex')
add('h','A Escopo e rastreabilidade')
add('p',f'O HEAD inicial foi confirmado em {meta["base_commit"]}, exatamente a base informada. Não havia commits posteriores para comparar. A árvore tinha somente .worktrees/ não rastreada, preservada. Não foi encontrado AGENTS.md no repositório ou seus diretórios ancestrais. Branch local: codex/exporural-2028-revisao. Commit final do código e dos dados desta implementação: {commit}. Este relatório é empacotado em commit documental posterior; o hash desse empacotamento é informado na entrega e pode ser conferido por git log. Não houve push, PR ou publicação.')
add('p',f'Revisão localizada: {revision}; versão geométrica da proposta: 6. A referência anterior do repositório continua em 2026.4-exporural.1, versão 5. OFFICIAL_REFERENCE_REVISION permanece 2026.4. O contador de versão de cada geometria persistida deverá avançar a partir da versão encontrada no banco, não ser sobrescrito com 6. Os nomes históricos 2026 e os demais setores não foram renomeados.')
add('p','Foram alterados apenas lotes R/S, sete vias locais e a apresentação necessária à revisão. Os polígonos, identificadores e áreas das demais entidades são preservados por identidade na fixture e verificados em teste. Pavilhões, estacionamento, vegetação, apoio, câmera global, permissões e regras comerciais mantêm seus dados. Nenhuma ferramenta de teste voltou à interface pública. A única entrada de QA fica em scripts/exporural/preview.html, fora do build público.')
add('h','B Fontes e inventário proposto')
add('p','Os arquivos presentes diferem da dimensão declarada no briefing: ambos os PNG têm 9934 × 7017 pixels nativos, com metadado aproximado de 300 dpi. O visual anexado ao chat foi reduzido para exibição. DPI não foi usado como escala métrica. Não havia PDF ou CAD original; nenhum arquivo desse tipo foi lido. As cópias exatas estão em fontes/.')
for src in meta['sources']:
    add('p',f'{src["filename"]} — {src["width"]} × {src["height"]} px. SHA-256: {src["sha256"]}.')
add('p','Fenasoja_Parque_Ajustes_300dpi.png orienta áreas, cotas e limites; Fenasoja_So_Numeros_300dpi.png permite extrair as faces fechadas e confirmar sequência/vizinhança. A fotografia WhatsApp serve somente como antes visual, sem comprovar situação comercial ou identidade cadastral. Os pontos vermelhos/verdes indicam infraestrutura e não foram convertidos em status de venda.')
table(['Quadra','Antes no repositório','Proposta','Variação'],[
 ['R','59 lotes · 31.492,94 m²','65 lotes · 29.564,26 m²','+6 lotes · −1.928,68 m²'],
 ['S','36 lotes · 16.391,93 m²','35 lotes · 16.203,53 m²','−1 lote · −188,40 m²'],
 ['Total','95 lotes · 47.884,87 m²','100 lotes · 45.767,79 m²','+5 lotes · −2.117,08 m²']])
add('p','A redução total não representa automaticamente a área da rua. A parcela sem número de 568,78 m² fica excluída do inventário comercial, sem código R-66, sem preço e com destinação pendente. Q-S-36 não integra o inventário proposto. A presença de R-60 a R-65 não define quantos UUIDs criar. Essa decisão depende da linhagem física.')
add('h','C Manifesto e correspondência física')
add('p','manifesto_lotes.json contém 100 entradas explícitas, com código, quadra, número, área decimal, revisão, fonte e hash. A soma foi conferida com Decimal no gerador e inteiros de centésimos nos testes. manifesto_lotes.csv, fixtures, exportações e tabelas deste relatório derivam desse mesmo manifesto. official_area_sqm é a transcrição proposta; calculated_area_sqm vem do anel e continua separado.')
byblock={b:[m for m in manifest if m['block']==b] for b in ['R','S']}
table(['Código R','Área oficial m²','Código S','Área oficial m²'],[
 [r['public_identifier'],fmt(r['official_area_sqm']),byblock['S'][i]['public_identifier'] if i<35 else '',fmt(byblock['S'][i]['official_area_sqm']) if i<35 else '']
 for i,r in enumerate(byblock['R'])])
add('p',f'crosswalk_linhagem.json e .csv organizam {len(cross)} grupos físicos, cobrindo cada um dos 95 códigos anteriores e cada um dos 100 propostos uma vez. Registram revisões, áreas, hashes, versões, evidência, confiança, ação candidata e aprovações. Todos permanecem PENDING_APPROVAL e todos os UUIDs reais permanecem nulos. Nem códigos iguais nem este alinhamento raster comprovam continuidade jurídica. A tabela abaixo resume todos os grupos; os arquivos preservam o detalhe por código.')
actions={'parcela_preservada_candidata':'Preservação candidata','renumeracao_candidata':'Renumeração candidata',
 'ajuste_area_limite':'Ajuste de área e limite','divisao_candidata':'Divisão candidata',
 'reparcelamento_pendente':'Reparcelamento pendente','correspondencia_pendente':'Correspondência pendente'}
table(['Grupo físico','Códigos anteriores','Códigos propostos','Ação candidata'],[
 [c['physical_parcel'],codes(c['previous_codes']),codes(c['proposed_codes']),actions.get(c['action'],c['action'].replace('_',' '))]
 for c in cross])
add('p','O caso R-56 foi tratado explicitamente: o antigo lote de 471 m² está no conjunto inferior e é candidato espacial ao novo R-62. O novo R-56 tem 249,03 m² e pertence à subdivisão da antiga ilha R-41 a R-43. Não há UPDATE, associação comercial ou revisão de preço gerados por igualdade nominal de R-56.')
add('h','D Geometria e circulação')
add('p','O gerador extrai contornos fechados da imagem nativa, simplifica em 3,2 pixels nativos, aproxima pontas compartilhadas até 1,4 pixels da grade de inspeção e insere junções coerentes. As formas não foram esticadas para atingir áreas oficiais. Cantos radiados, trapézios e o limite inferior inclinado de R-62 a R-65 permanecem representados. Cada anel local é fechado, finito, simples e orientado de forma consistente. O ponto de rótulo é interno à parcela, calculado por representative_point, não pela média dos vértices.')
add('p','As coordenadas utilizam LOCAL_NORMALIZED, SRID 0, em unidades de cena; 0,15 unidade de cena corresponde a um metro na calibração histórica. Não são longitude/latitude ou WGS84. A origem raster está no canto superior esquerdo, X para a direita e Y para baixo. A grade de inspeção tem largura 1888 e fator nativo 5,2616525424. Aplica-se primeiro a transformação afim para o referencial histórico; somente depois ocorre a conversão para coordenadas locais. A calibração global não foi alterada.')
add('p','Com u e v na grade de inspeção: sourceX = 1,5056770501648136u − 0,00038893441509291193v + 3151,3107458024915; sourceY = 0,017094028627930902u + 1,5053614459728695v + 770,1650528245467. LocalX = (sourceX − 600) × 120 / 5500 − 60; LocalY = (sourceY − 900) × 120 / 5500 − 45,2727272727. Rotação do eixo X: 0,650454 grau. Os coeficientes completos estão em calibracao.json.')
add('p',f'Foram usados seis controles distribuídos não colineares e três checagens independentes, todos referenciados à geometria do repositório, sem levantamento de campo. RMSE dos controles: {cal["controlsRmseMeters"]:.3f} m; máximo: {cal["controlsMaxMeters"]:.3f} m. RMSE das checagens: {cal["independentChecksRmseMeters"]:.3f} m; máximo: {cal["independentChecksMaxMeters"]:.3f} m. A tabela de controles, observações e resíduos está em calibracao.json. Não se declara precisão centimétrica.')
add('p','A tolerância antiga de 0,15% não foi relaxada: os 100 polígonos a excedem, com desvio absoluto máximo de 7,657079% entre área calculada e oficial. A combinação de raster e registro histórico não comprova metragens geométricas cadastrais. A entrega conserva UNVALIDATED/NEEDS_REVIEW. Para validação métrica definitiva são necessários controle mais preciso e/ou arquivo técnico original confirmado; corrigir limites exige gerar nova revisão e repetir QA, sem ajustar área por escala X.')
add('p','Há também conflito físico preexistente com os apoios preservados: no repositório antigo, B37 sobrepõe S-24 em 156,393 m² e B38 sobrepõe S-25 em 152,669 m². No registro novo, B37 cruza S-22 (1,029 m²) e S-23 (154,783 m²); B38 cruza S-23 (1,312 m²) e S-24 (150,782 m²). São áreas calculadas na calibração histórica, sem precisão cadastral. B37/B38 não foram movidas ou apagadas, nem os contornos da planta recortados para esconder o problema. C4, B7, B8 e D3 não apresentam invasão pelos novos lotes. evidencias/conflitos_apoios.json e geometria_qa.json registram o bloqueio. O SQL aborta essas interseções mesmo com aprovações preenchidas; resolver exige decisão cadastral explícita e uma revisão compatível com os limites autorizados. A aceitação geométrica completa ainda NÃO foi atingida.')
add('p','A transversal EXPORURAL-ACESSO-TRANSVERSAL-01 começa na Bruno Schwartz, passa entre S-04/S-05, cruza Johan Muller, passa entre R-31/R-32, cruza Gustavo Bessel e termina entre R-22/R-23 no limite indicado. Não atravessa a faixa superior S, o estacionamento ou o Mirante. A largura documental nominal é 6,00 m; as concordâncias raster estão representadas, mas não certificadas em campo. Nome mostrado: Via interna — denominação a confirmar.')
add('p','Rua 15 de Novembro fica entre R-39/R-40 e à esquerda de R-48/R-55, separando esse conjunto de R-28. A passagem entre R-50/R-51 e R-57/R-58 recebe outro ID técnico. As sete divisas horizontais 48/55 a 54/61 são bordas de lotes e não entidades de rua. O corredor existente entre R-25/R-26 também tem identidade própria. As duas passagens já desenhadas na planta, anteriormente sem entidade individual, explicam duas das três entidades viárias acrescentadas.')
table(['Identificador técnico','Nome exibido','Tratamento'],[
 [r['public_identifier'],r['name'],'Nova transversal' if 'ACESSO-TRANSVERSAL' in r['public_identifier'] else 'Passagem existente individualizada' if 'PASSAGEM-INTERNA' in r['public_identifier'] else 'Via existente ajustada localmente'] for r in roads])
add('p','As vias existentes conservam a identidade persistida quando o preflight a confirmar. RUA-UBIRETAMA-LATERAL-R55 e RUA-LESTE-EXPORURAL foram auditadas, preservadas e incluídas na checagem de sobreposição contra os lotes novos. A menção histórica a R55 não é prova do lote atual: atualizar descrição somente mediante confirmação, sem recriar a via ou trocar seu UUID. Interseções rua-rua são permitidas; lote-rua não.')
add('h','E Código e consumidores auditados')
add('p','exporuralReference2028.ts constrói uma fixture explícita e rejeita entrada source=database ou projeto com orgId persistido. Seus IDs incluem reference:preview e a revisão; não servem como IDs de banco. Os 100 lotes locais são BLOCKED, NOT_FOR_SALE, sem preço, comprador ou vínculo inventado. A entrada de preview substitui o cliente Supabase por um stub que lança erro para qualquer RPC, tabela ou escrita. O CSP limita conexões ao servidor local.')
add('p','O preview utiliza CommercialMapShell, CommercialMapPage, o mesmo CommercialMapCanvas, controles e renderer. Não existe mapa paralelo. Lista, busca, filtros, exportação e dashboard recebem as entidades/lotes revisados e a área do manifesto. A renderização dos números reutiliza o atlas instanciado existente, ativa-se somente por metadata.geometryRevision de Exporural 2028 e usa o mesmo labelAnchor. Picking, contorno e cartão continuam associados ao mesmo entityId. A apresentação SOLD vermelho com cadeado não foi alterada. alteracoes_codigo.patch contém as mudanças de frontend, scripts e testes para o Lovable conferir contra o HEAD atual e integrar preservando eventuais melhorias posteriores.')
add('p','Duas adaptações compartilhadas eram necessárias: o paisagismo de R-13/R-14/R-02 passou a localizar cantos por geometria em vez de índices fixos, eliminando uma degeneração com anéis que têm novas junções; e a seleção/carrinho descarta entidades R/S cujo UUID, código, área, revisão ou geometria mudou entre leituras. Essa limpeza fecha checkout obsoleto sem cancelar reserva ou editar contrato. A numeração usa a camada já existente sem duplicá-la na visão pública.')
add('p','officialReference2026.ts e seus expectedLotCounts R59/S36 continuam representando a revisão antiga usada pelo bootstrap histórico. As contagens 65/35 e 116/100 pertencem à revisão 2028, seus manifestos e validadores. reconcileExporuralReference.ts permanece autoritativo para o banco, sem sobrepor a proposta local. commercialMapService, proteções de useCommercialMap, preços e regras de acesso não foram removidos. Segmentação usa a área Exporural da própria entidade; a rua nova participa da circulação, do segmento e das entidades apresentadas. A navegação de visita continua derivando superfícies das entidades ROAD.')
add('h','F Banco e preflight obrigatório')
add('p','Não houve conexão com Supabase. O esquema foi lido das migrations versionadas, especialmente 20260710010000_create_commercial_map.sql, 20260804090000_create_commission_map_segments.sql e 20260811153000_apply_exporural_reference_2026_4_fidelity.sql e suas duplicatas cronológicas. O ambiente vivo pode divergir; os nomes e constraints devem ser confirmados pelo Lovable antes de qualquer aplicação.')
table(['Tabela','Campos e finalidade confirmados'],[
 ['map_projects','id, org_id, reference_revision, active_version, is_published, is_archived; controle do projeto e publicação'],
 ['map_entities','id, project_id, layer_id, parent_entity_id, public_identifier, classification, verification_status, is_sellable, is_archived, metadata, segment_id'],
 ['map_entity_geometries','id, entity_id, geometry JSONB, native_geometry Polygon SRID 0 gerada, version, is_current, calibration_version, elevation, extrusion_height, rotation'],
 ['map_geometry_versions','geometry_id, entity_id, version e geometria anterior; histórico versionado'],
 ['commercial_lots','id, entity_id, public_identifier, block, lot_number, official_area_sqm, calculated_area_sqm, status, area_validation_status, archived_at'],
 ['map_segments','id, project_id, slug, source_reference, boundary_data, is_active; expectedEntityCount, expectedLotCount, lineageBaselineAt'],
 ['map_lot_lineage','source_lot_id, target_lot_id, relationship, created_by; SPLIT_FROM, MERGED_FROM, SUPERSEDES'],
 ['map_reference_migration_snapshots','project_id, area_code, source_revision, payload_hash, status, snapshot, apply_result e dados de rollback'],
 ['map_activity_logs','org_id, project_id, action, reason, before_state, after_state, actor_user_id'],
 ['lot_prices','lot_id, pricing_mode, base_price, price_per_sqm, asking_price, minimum_price, is_active; sem alterações automáticas'],
 ['lot_reservations e lot_negotiations','lot_id, status e histórico de negociações/reservas; rechecagem sob bloqueio'],
 ['lot_sales, lot_contracts e lot_contract_versions','lot_id ou contract_id, valores e snapshots históricos; não reassociar nem sobrescrever']])
add('p','Executar primeiro somente preflight.sql, em transação de leitura repetível e encerrada com ROLLBACK. Definir os UUIDs reais de org_id e project_id na sessão, confirmar projeto, ator autorizado, segmento, revisões, calibração, inventário, unicidade inclusive arquivados, versões/hash MD5 da geometria corrente, updated_at, contratos, reservas, vendas, negociações, preços, histórico e linhagem. Exportar os resultados e confrontar parcela física por parcela física com o crosswalk. Divergência exige abortar e relatar, nunca apagar dados para ajustar contagens.')
add('p','A fixture corrente tem 108 entidades/95 lotes; a migration histórica citada esperava 111/95. Não se presume qual está implantada. O preflight deve explicar apoios e vias realmente presentes. Na fixture proposta são 116/100: 100 lotes + 10 vias + 3 entidades de área/quadra + 3 apoios correntes. Uma base viva de 111/95 com os mesmos 102 registros afetados resultaria em 119/100, não 116/100. O SQL conta ambos separadamente e exige expected_entity_count aprovado.')
add('p','O script exige auth.uid() auditável, map.admin e os privilégios de manutenção necessários. Na migration consultada, authenticated tem apenas SELECT em map_reference_migration_snapshots. O Lovable deve confirmar o executor de manutenção já autorizado do projeto; não desativar RLS, não criar bypass aberto nem remover verificações. Se o executor não estiver disponível, interromper antes da proposta transacional.')
add('h','G Linhagem e preservação comercial')
add('p','approvals_resolvidas.json é o formulário completo de aplicação, ainda NÃO RESOLVIDO: contém 102 entidades anteriores (95 lotes e sete vias), 110 destinos (100 lotes e dez vias), versões/hash/timestamps esperados e alocações pai-filho. null e false são bloqueios intencionais. Para continuidade comprovada, preencher source_entity_id/source_lot_id com UUIDs reais e action=preserve. Para parcela resultante nova, action=create somente após aprovar pais, filhos e destinação do pai. Novos UUIDs são gerados pelo PostgreSQL na transação autorizada, nunca fornecidos pela fixture.')
add('p','Todas as aprovações devem citar evidência e responsável. Pais com venda ou contrato histórico são bloqueados pelo script genérico; exigem procedimento específico autorizado. Pai com reserva/negociação ativa não pode ser arquivado automaticamente. Uma continuidade com vínculo ativo exige resolução explícita documentada. Não copiar venda para vários filhos, não cancelar reservas, não apagar contratos, não transportar negócio por renumeração e não preencher preços por suposição.')
add('p','Nos registros preservados, o SQL conserva status, identidade, linhas de preços e vínculos. Novos lotes são BLOCKED e is_sellable=false, sem inserir preço zero ou linha de preço. NEEDS_REVIEW/UNVALIDATED não deve ser convertido em VERIFIED/VALIDATED para contornar a revisão. A alteração de área pode afetar uma cotação futura por m²; a política e os valores futuros devem ser aprovados separadamente. Snapshots e valores de negócios firmados não são recalculados.')
add('h','H Aplicação transacional proposta')
add('p','apply_exporural_reference_2026 foi inspecionada: recebe arrays e reconcilia registros por public_identifier, pressupondo a antiga composição. Não é segura para esta renumeração e não deve ser reutilizada. sync_commercial_map_reference_2026 também não deve ser usado; ele poderia repovoar o parque inteiro. Os SQL desta entrega estão somente na pasta documental, fora de supabase/migrations e sem execução automática no carregamento do mapa.')
add('p','Após resolver e autorizar cartografia/linhagem, preparar carregar_payload.sql com o JSON aprovado, abrir BEGIN na mesma sessão, carregar a tabela temporária, executar migration_proposta.sql e verificacao_pos_migracao.sql. Só fazer COMMIT após comparar resultados e snapshots. O script usa bloqueio do projeto, advisory lock, bloqueio dos registros comerciais, versões/timestamps/hash esperados, aprovação por destino, escopo R/S, topologia e colisões. RLS e autorização continuam ativas.')
add('p','Os códigos antigos são primeiro substituídos por identificadores temporários derivados de UUID dentro da transação; então se atribuem os códigos finais aos UUIDs aprovados. Pais retirados são arquivados com identificador de auditoria e histórico, sem DELETE físico. A geometria anterior é gravada em map_geometry_versions. O hash do payload, a revisão e o estado pós-aplicação são gravados no snapshot. Repetir o mesmo payload e estado retorna sem alterações; payload diferente ou deriva posterior aborta. As garantias transacionais ainda precisam ser executadas contra banco local isolado antes do banco vivo.')
add('p','O conjunto proposto deve avançar active_version e manter reference_revision global, atualizar source_reference do segmento e gravar expectedEntityCount e expectedLotCount com suas contagens reais independentes. lineageBaselineAt é redefinido após inserir a linhagem aprovada, evitando somar a mesma divisão novamente. A entidade de rua aumenta entidades, não lotes. O projeto fica is_published=false até a validação/publicação conforme o processo existente.')
add('h','I Snapshot e restauração')
add('p','O snapshot inclui projeto, segmento, entidades, geometrias, lotes, preços, reservas, negociações, vendas, contratos, versões contratuais e linhagem afetada. O hash externo do estado fora do escopo é comparado antes/depois. rollback_por_snapshot.md fornece o procedimento condicionado: bloquear, verificar que não houve venda ou mudança posterior, preservar uma cópia do estado atual e restaurar por UUID sob transação. Se houver deriva, abortar a restauração e fazer reconciliação específica; nunca restaurar por cima de negócio novo. Filhos ficam arquivados e a geometria restaurada avança o histórico. Nenhum rollback foi executado.')
add('h','J Validação publicação e cache')
add('p','A releitura após COMMIT deve ocorrer em sessão nova. Conferir as 100 áreas individuais com o manifesto, não só totais; revisar código/entity_id/lot_id, labelAnchor, polígonos, R-56=249,03, via transversal e separação da Rua 15 de Novembro. Conferir contratos/preços/reservas contra snapshot. Validar get_commission_map_segment_inventory, baselineEntityCount, baselineLotCount e lineageDelta em mapa completo e por comissão. Invalidar consultas React Query do mapa, comissão, lotes/preços e dashboard; recarregar a partir do servidor e limpar seleção antiga. Publicar apenas conforme os bloqueios de revisão do projeto permitirem.')
add('p','Os testes locais cobriram códigos/somas, sequências e vizinhança, rótulos internos, polígonos simples/fechados, divisas compartilhadas, ausência de lote-lote e lote-via em todas as ruas carregadas, B7/B8/D3 preservados, ausência de S-36/R-66, identidade R-56, fixture determinística, seleção obsoleta, métricas e invariância externa. O teste de paisagismo da referência antiga também passou. A proposta SQL passou pelo parser PostgreSQL/PLpgSQL pglast; esse resultado é somente sintático, não execução nem prova de RLS, triggers ou concorrência.')
add('p','Comandos e saídas: npm run typecheck e ESLint dos arquivos alterados terminaram com código 0; npm run build -- --manifest passou, com avisos existentes de tamanho de chunks e Browserslist antigo. A checagem focal inicial passou 20 testes de revisão/paisagismo; ao ampliar a proteção para B37/B38/C4, a suíte manteve UMA FALHA de aceitação por sobreposição com B37/B38. Essa verificação não foi removida ou relaxada para passar. Os demais testes de revisão, paisagismo, política pública e SOLD passaram nas execuções registradas. Foi necessário aumentar somente o timeout do runner para 60 s em testes pesados de paisagismo; nenhuma tolerância geométrica foi relaxada. A suíte ampliada anterior registrou 59 passes e três falhas: uma por timeout depois resolvida, uma expectativa antiga do explorador (262 versus 264) e um hash histórico de arquitetura. A auditoria inicial também encontrou expectativa global 1577 versus 1579. As falhas de snapshots/contagens históricas foram preservadas e não são consideradas passes.')
add('p','evidencias/runtime contém capturas reais de antes/depois na mesma câmera, viewport desktop 1440 × 1000, zoom e camadas para seis vistas: geral, faixas S, faixa central R, subdivisões, perímetro 62–65 e transversal. O clique no novo R-56 apresentou 249,03 m²; trocar revisão limpou seleção. Foi injetada perda de contexto WebGL e verificada recuperação. Viewport mobile 390 × 844 e a rota completa com fallback acessível também foram exercitados. resultados.json registra navegador, erros, bloqueios de rede, estado, inventário e recursos. Pranchas técnicas de sobreposição em evidencias/pranchas usam exatamente as coordenadas da fixture, mas não são screenshots do renderer.')
add('p','A recuperação WebGL teve resultados mistos nas repetições: execucao-validada-anterior.json registra sucesso com as mesmas identidades de Canvas, renderer, câmera e controles; tentativa-contexto-timeout.json registra uma espera de recuperação que excedeu 90 s durante execução concorrente de outras tarefas locais. Isso não comprova a causa da falha. O ensaio final fica em resultados.json. Estabilidade sob pressão de recursos permanece pendente e não se declara aceite universal do renderer.')
add('p','Limites: as rotas reais de componentes foram testadas com fixtures e Supabase substituído por stub, sem sessão autenticada real nem dados vivos de comissão. O fallback acessível denominado 2D é a lista/tabela existente; não foi criada uma segunda maquete 2D. Não houve teste em celular físico, Safari/iOS, Android, produção, CI remota ou banco isolado, pois Docker/psql não estavam disponíveis. Não há certificação universal de FPS/memória ou prazo de cold start. A primeira captura revelou e motivou a correção dos índices do paisagismo; uma navegação de fallback excedeu timeout sob carga e foi repetida no ensaio final, concluído sem erros. Os logs registram as tentativas de diagnóstico e o JSON de runtime registra a execução final.')
add('h','K Pendências para autorização')
table(['Pendência','Responsável pela resolução','Efeito'],[
 ['Calibração e diferença de área calculada','Responsável cartográfico e Lovable','Bloqueia aceite cadastral e aplicação proposta'],
 ['B37 e B38 sobrepostos às parcelas S','Responsável cadastral e organização','Bloqueio geométrico explícito; preservar estruturas até decisão autorizada'],
 ['UUIDs e identidade de cada parcela','Gestão cadastral/comercial e Lovable','Bloqueia renumeração e reassociação'],
 ['Divisões e reparcelamentos ambíguos','Gestão cadastral/comercial','Aprovar cada pai/filho e arquivamento'],
 ['Reservas vendas e contratos ativos ou históricos','Responsável comercial/jurídico','Procedimento explícito sem transferência automática'],
 ['Nome da transversal e passagens','Organização da feira','Manter denominação a confirmar'],
 ['Destinação dos 568,78 m² sem número','Organização da feira','Excluir de venda e do inventário de 100 lotes'],
 ['Cadastro preço e autorização de novas parcelas','Responsável comercial','Manter BLOCKED e sem preço até liberação'],
 ['Preflight banco isolado executor e publicação','Lovable','Não aplicar por simples build ou commit']])
add('h','L Arquivos e ordem de uso')
table(['Arquivo ou pasta','Utilização'],[
 ['RELATORIO_PARA_LOVABLE.md e .docx','Mesmo conteúdo preenchido desta entrega'],
 ['alteracoes_codigo.patch','Frontend, scripts e testes; comparar com a base e integrar sem reverter mudanças posteriores'],
 ['manifesto_lotes.json e .csv','100 parcelas e áreas oficiais propostas'],
 ['crosswalk_linhagem.json e .csv','Grupos físicos e todas as correspondências pendentes'],
 ['geometrias_lotes.json e geometria_vias.json','Anéis usados no preview, coordenadas locais, âncoras e SHA-256'],
 ['calibracao.json e fontes_e_revisao.json','Transformação, controles, resíduos e origem'],
 ['referencia_anterior.json','Referência do repositório; não é snapshot de banco'],
 ['payload_preview.json','Fixture sintética para inspeção; NÃO enviar a RPC de persistência'],
 ['approvals_resolvidas.json e carregar_payload.sql','Plano explicitamente bloqueado; completar com preflight e aprovações'],
 ['preflight.sql','Executar primeiro e somente em leitura'],
 ['migration_proposta.sql','Aplicar só após autorizações e ensaio isolado, na mesma transação do payload'],
 ['verificacao_pos_migracao.sql','Verificar antes do COMMIT e em nova sessão depois'],
 ['rollback_por_snapshot.md','Restauração condicionada à ausência de mudanças posteriores'],
 ['REPRODUCAO.md e TESTES_DE_BANCO_PENDENTES.md','Comandos locais e matriz de persistência ainda não executada'],
 ['evidencias/ e fontes/','Logs, screenshots, sobreposições e fontes exatas'],
 ['arquivos_alterados.txt e inventario_arquivos_sha256.json','Arquivos alterados e integridade do pacote'],
 ['PROMPT_PARA_LOVABLE.txt','Texto pronto para encaminhar com este pacote']])
add('p','Estado da entrega: análise do repositório concluída; código e preview implementados; testes locais executados com os limites registrados; pendências cartográficas e de linhagem abertas; pacote de banco preparado e não aplicado; aplicação Supabase e publicação pendentes do Lovable. Nenhuma alteração foi aplicada ao Supabase pelo Codex')

# One content model, two output formats. No separate hand-edited Word narrative.
md=[]
doc=Document();section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.7)
section.left_margin=section.right_margin=Inches(.7)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
for style,size in [('Title',22),('Heading 1',15)]:
    doc.styles[style].font.name='Arial';doc.styles[style].font.size=Pt(size);doc.styles[style].font.color.rgb=RGBColor(0,0,0)
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(15)
footer=section.footer.paragraphs[0];footer.alignment=2
footer.add_run('Exporural 2028  •  Proposta local  •  ').font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for kind,value in blocks:
    if kind=='table':
        md.append('\n'.join(['| '+' | '.join(row)+' |' for row in [value[0],['---']*len(value[0]),*value[1:]]]))
        t=doc.add_table(rows=1,cols=len(value[0]));t.autofit=False
        widths=[7.1/len(value[0])]*len(value[0])
        if len(value[0])==2:widths=[2.35,4.75]
        if value[0][0]=='Grupo físico':widths=[2.2,1.55,1.7,1.65]
        table_width=t._tbl.tblPr.find(qn('w:tblW'))
        table_width.set(qn('w:type'),'dxa');table_width.set(qn('w:w'),str(round(7.1*1440)))
        for i,w in enumerate(widths):t.columns[i].width=Inches(w)
        for i,txt in enumerate(value[0]):t.rows[0].cells[i].text=txt
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in value[1:]:
            cells=t.add_row().cells
            for i,txt in enumerate(row):cells[i].text=txt
        for ri,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            for i,cell in enumerate(row.cells):
                cell.width=Inches(widths[i])
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                    for run in p.runs:run.font.size=Pt(9.5);run.bold=ri==0
                if ri==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8ECEA');cell._tc.get_or_add_tcPr().append(shade)
        doc.add_paragraph().paragraph_format.space_after=Pt(1)
    else:
        md.append(('# ' if kind=='title' else '## ' if kind=='h' else '')+value)
        p=doc.add_paragraph(value,'Title' if kind=='title' else 'Heading 1' if kind=='h' else 'Normal')
        if kind=='p':p.paragraph_format.widow_control=True
(root/'RELATORIO_PARA_LOVABLE.md').write_text('\n\n'.join(md)+'\n',encoding='utf-8')
doc.core_properties.title='Revisão cadastral e geométrica da Exporural FENASOJA 2028'
doc.core_properties.author='Codex'
doc.core_properties.subject='Entrega local para revisão e aplicação autorizada pelo Lovable'
doc.save(root/'RELATORIO_PARA_LOVABLE.docx')
# Exact textual equivalence including all table cells, before layout rendering.
from docx.table import Table
from docx.text.paragraph import Paragraph
observed=[]
for child in doc.element.body:
    if child.tag==qn('w:p'):
        text=Paragraph(child,doc).text
        if text:observed.append(text)
    elif child.tag==qn('w:tbl'):
        for row in Table(child,doc).rows:
            observed.extend(c.text for c in row.cells)
expected=[]
for kind,value in blocks:
    if kind=='table':expected.extend(cell for row in value for cell in row)
    else:expected.append(value)
assert observed==expected,'MD/DOCX content mismatch'
(root/'evidencias/documento-conteudo.json').write_text(json.dumps({'same_content':True,'content_blocks':len(blocks),'code_commit':commit},indent=2)+'\n',encoding='utf-8')
print('Report generated:',len(blocks),'content blocks; identical Markdown/Word text')
